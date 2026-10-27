#!/usr/bin/env python3
"""
Universal Client-Ready Audit Report Engine (v2.1 - Sanitized & Hardened)
Generates standardized "Issues-Only" Executive DOCX and PDF Reports
with full XML borders, cell shading, explicit column widths, and Playwright Chromium.
Properly escapes HTML entities to prevent unescaped markup (e.g. <h1> tags in text)
from breaking layout.
"""

import argparse
import html
import json
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from playwright.sync_api import sync_playwright

HEX_NAVY = "0A192F"
HEX_BLUE = "2563EB"
HEX_RED = "DC2626"
HEX_RED_BG = "FEF2F2"
HEX_BLUE_BG = "EFF6FF"
HEX_GRAY_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_DARK = "0F172A"
HEX_MUTED = "64748B"

RGB_NAVY = RGBColor(10, 25, 47)
RGB_BLUE = RGBColor(37, 99, 235)
RGB_RED = RGBColor(220, 38, 38)
RGB_DARK = RGBColor(15, 23, 42)
RGB_MUTED = RGBColor(100, 116, 139)
RGB_WHITE = RGBColor(255, 255, 255)


def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_border(cell, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, color in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if color:
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))


def build_docx_report(data, output_path):
    doc = docx.Document()

    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Running Header & Footer
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"{data.get('client_name', 'CLIENT').upper()}  |  TECHNICAL SEO & ON-PAGE AUDIT")
    hrun.font.name = "Arial"
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGB_MUTED

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun = fp.add_run(f"{data.get('client_name', 'Client')} - SEO Audit Report  |  Confidential")
    frun.font.name = "Arial"
    frun.font.size = Pt(8)
    frun.font.color.rgb = RGB_MUTED

    # Title & Subtitle
    badge_p = doc.add_paragraph()
    badge_p.paragraph_format.space_before = Pt(0)
    badge_p.paragraph_format.space_after = Pt(2)
    r_badge = badge_p.add_run(data.get("badge_text", "B2B TECHNICAL SEO AUDIT").upper())
    r_badge.font.name = "Arial"
    r_badge.font.size = Pt(9.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGB_BLUE

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(2)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run(data.get("report_title", "TECHNICAL SEO & ON-PAGE AUDIT REPORT"))
    r_title.font.name = "Arial"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGB_NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    r_sub = sub_p.add_run(f"{data.get('client_name', 'Client')} ({data.get('domain', '')}) — Issues & Gap Analysis")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGB_MUTED

    # Executive Health Score & Metadata Box
    score_box_table = doc.add_table(rows=1, cols=2)
    score_box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    score_box_table.autofit = False

    c_score = score_box_table.rows[0].cells[0]
    c_meta = score_box_table.rows[0].cells[1]
    c_score.width = Inches(2.1)
    c_meta.width = Inches(4.6)

    set_cell_background(c_score, HEX_RED_BG)
    set_cell_border(c_score, top=HEX_RED, bottom=HEX_RED, left=HEX_RED, right=HEX_RED, sz="12")
    set_cell_margins(c_score, top=160, bottom=160, left=160, right=160)

    p_sc_val = c_score.paragraphs[0]
    p_sc_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sc_val.paragraph_format.space_after = Pt(0)
    r_val = p_sc_val.add_run(f"{data.get('health_score', 50)} / 100")
    r_val.font.name = "Arial"
    r_val.font.size = Pt(28)
    r_val.font.bold = True
    r_val.font.color.rgb = RGB_RED

    p_sc_lbl = c_score.add_paragraph()
    p_sc_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sc_lbl.paragraph_format.space_before = Pt(2)
    p_sc_lbl.paragraph_format.space_after = Pt(0)
    r_lbl = p_sc_lbl.add_run("OVERALL HEALTH SCORE\nCritical Remediation Required")
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(8)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = RGB_RED

    set_cell_background(c_meta, HEX_GRAY_BG)
    set_cell_border(c_meta, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER, sz="4")
    set_cell_margins(c_meta, top=140, bottom=140, left=160, right=160)

    meta_items = data.get("meta_items", [])
    p_m = c_meta.paragraphs[0]
    p_m.paragraph_format.space_after = Pt(2)
    p_m.paragraph_format.line_spacing = 1.15

    for idx, (lbl, val) in enumerate(meta_items[:5]):
        if idx > 0:
            p_m = c_meta.add_paragraph()
            p_m.paragraph_format.space_after = Pt(2)
            p_m.paragraph_format.line_spacing = 1.15
        r_l = p_m.add_run(f"{lbl} ")
        r_l.font.name = "Arial"
        r_l.font.size = Pt(8.5)
        r_l.font.bold = True
        r_l.font.color.rgb = RGB_DARK

        r_v = p_m.add_run(val)
        r_v.font.name = "Arial"
        r_v.font.size = Pt(8.5)
        r_v.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def add_h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGB_NAVY
        return p

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(9)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGB_DARK
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGB_DARK
        return p

    # 1. Executive Summary
    add_h("1. Executive Summary & Verified Tool Baseline")
    for para in data.get("executive_summary", []):
        add_p(para.get("text", ""), para.get("bold_prefix", ""))

    # 2. Scorecard Table
    add_h(f"2. Website Technical Health Scorecard ({data.get('health_score', 50)} / 100)")
    sc_table = doc.add_table(rows=1, cols=5)
    sc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sc_table.autofit = False

    col_widths = [Inches(1.8), Inches(0.6), Inches(0.8), Inches(0.8), Inches(2.7)]
    headers = ["Audit Category", "Weight", "Score", "Status", "Primary Observed Defect"]
    hdr_row = sc_table.rows[0]
    prevent_row_split(hdr_row)
    set_repeat_table_header(hdr_row)

    for idx, h in enumerate(headers):
        c = hdr_row.cells[idx]
        c.width = col_widths[idx]
        c.text = h
        set_cell_background(c, HEX_NAVY)
        set_cell_margins(c, top=140, bottom=140, left=100, right=100)
        set_cell_border(c, top=HEX_NAVY, bottom=HEX_NAVY, left=HEX_NAVY, right=HEX_NAVY)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [1, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.runs[0]
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGB_WHITE

    for row_idx, r_data in enumerate(data.get("scorecard", [])):
        row = sc_table.add_row()
        prevent_row_split(row)
        bg_color = HEX_GRAY_BG if row_idx % 2 == 1 else "FFFFFF"
        is_total = "TOTAL" in str(r_data[0])
        if is_total:
            bg_color = HEX_BLUE_BG

        for c_idx, val in enumerate(r_data):
            c = row.cells[c_idx]
            c.width = col_widths[c_idx]
            c.text = str(val)
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=100, bottom=100, left=100, right=100)
            set_cell_border(c, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [1, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(8)

            if c_idx == 3:
                r.font.bold = True
                if val in ["Critical", "Fail"]:
                    r.font.color.rgb = RGB_RED
                elif val == "Pass":
                    r.font.color.rgb = RGBColor(16, 128, 64)
                elif val == "Warning":
                    r.font.color.rgb = RGBColor(200, 120, 0)
            elif is_total:
                r.font.bold = True
                if c_idx == 2:
                    r.font.color.rgb = RGB_RED

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. Issue Catalog
    issues = data.get("issues", [])
    add_h(f"3. Comprehensive Issue Catalog ({len(issues)} Verified Issues - Issues ONLY)")

    iss_table = doc.add_table(rows=1, cols=3)
    iss_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    iss_table.autofit = False

    iss_col_widths = [Inches(2.0), Inches(1.6), Inches(3.1)]
    iss_headers = ["Issue Title & Severity", "Affected URL(s)", "Observed Defect & Impact"]

    hdr_iss = iss_table.rows[0]
    prevent_row_split(hdr_iss)
    set_repeat_table_header(hdr_iss)

    for idx, h in enumerate(iss_headers):
        c = hdr_iss.cells[idx]
        c.width = iss_col_widths[idx]
        c.text = h
        set_cell_background(c, HEX_NAVY)
        set_cell_margins(c, top=140, bottom=140, left=100, right=100)
        set_cell_border(c, top=HEX_NAVY, bottom=HEX_NAVY, left=HEX_NAVY, right=HEX_NAVY)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.runs[0]
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGB_WHITE

    for row_idx, iss in enumerate(issues, 1):
        row = iss_table.add_row()
        prevent_row_split(row)
        bg_color = HEX_GRAY_BG if row_idx % 2 == 1 else "FFFFFF"

        c0 = row.cells[0]
        c0.width = iss_col_widths[0]
        set_cell_background(c0, bg_color)
        set_cell_margins(c0, top=100, bottom=100, left=100, right=100)
        set_cell_border(c0, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.15
        r0_num = p0.add_run(f"{row_idx}. {iss.get('title', '')}\n")
        r0_num.font.name = "Arial"
        r0_num.font.size = Pt(8)
        r0_num.font.bold = True
        r0_num.font.color.rgb = RGB_DARK

        sev = iss.get("severity", "Medium")
        r0_sev = p0.add_run(f"[{sev.upper()}]")
        r0_sev.font.name = "Arial"
        r0_sev.font.size = Pt(7.5)
        r0_sev.font.bold = True
        r0_sev.font.color.rgb = RGB_RED if sev == "Critical" else (RGBColor(200, 100, 0) if sev == "High" else RGBColor(100, 116, 139))

        c1 = row.cells[1]
        c1.width = iss_col_widths[1]
        set_cell_background(c1, bg_color)
        set_cell_margins(c1, top=100, bottom=100, left=100, right=100)
        set_cell_border(c1, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(iss.get("urls", ""))
        r1.font.name = "Arial"
        r1.font.size = Pt(7.5)
        r1.font.color.rgb = RGB_MUTED

        c2 = row.cells[2]
        c2.width = iss_col_widths[2]
        set_cell_background(c2, bg_color)
        set_cell_margins(c2, top=100, bottom=100, left=100, right=100)
        set_cell_border(c2, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(iss.get("impact", ""))
        r2.font.name = "Arial"
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGB_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Scope of Improvement
    add_h("4. Scope of Improvement")
    for sc in data.get("scope_of_improvement", []):
        add_p(sc.get("text", ""), sc.get("bold_prefix", ""))

    # 5. Immediate Fixation Needed
    add_h(f"5. Immediate Fixation Needed (Emergency Priority - {len(data.get('immediate_fixations', []))} Items)")
    for idx, fix in enumerate(data.get("immediate_fixations", []), 1):
        add_p(fix, f"{idx}. ")

    # 6. Strategic Roadmap
    add_h("6. What Will Be the Next Step (Strategic 4-Phase Roadmap)")
    rm_table = doc.add_table(rows=1, cols=3)
    rm_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rm_table.autofit = False

    rm_col_widths = [Inches(1.6), Inches(3.3), Inches(1.8)]
    rm_headers = ["Phase", "Key Objectives & Focus Areas", "Target Milestone"]

    hdr_rm = rm_table.rows[0]
    prevent_row_split(hdr_rm)
    set_repeat_table_header(hdr_rm)

    for idx, h in enumerate(rm_headers):
        c = hdr_rm.cells[idx]
        c.width = rm_col_widths[idx]
        c.text = h
        set_cell_background(c, HEX_NAVY)
        set_cell_margins(c, top=140, bottom=140, left=100, right=100)
        set_cell_border(c, top=HEX_NAVY, bottom=HEX_NAVY, left=HEX_NAVY, right=HEX_NAVY)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.runs[0]
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGB_WHITE

    for row_idx, rm in enumerate(data.get("roadmap", [])):
        row = rm_table.add_row()
        prevent_row_split(row)
        bg_color = HEX_GRAY_BG if row_idx % 2 == 1 else "FFFFFF"

        c0 = row.cells[0]
        c0.width = rm_col_widths[0]
        set_cell_background(c0, bg_color)
        set_cell_margins(c0, top=100, bottom=100, left=100, right=100)
        set_cell_border(c0, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(rm.get("phase", ""))
        r0.font.name = "Arial"
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = RGB_DARK

        c1 = row.cells[1]
        c1.width = rm_col_widths[1]
        set_cell_background(c1, bg_color)
        set_cell_margins(c1, top=100, bottom=100, left=100, right=100)
        set_cell_border(c1, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(rm.get("objectives", ""))
        r1.font.name = "Arial"
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGB_DARK

        c2 = row.cells[2]
        c2.width = rm_col_widths[2]
        set_cell_background(c2, bg_color)
        set_cell_margins(c2, top=100, bottom=100, left=100, right=100)
        set_cell_border(c2, top=HEX_BORDER, bottom=HEX_BORDER, left=HEX_BORDER, right=HEX_BORDER)
        p2 = c2.paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(rm.get("milestone", ""))
        r2.font.name = "Arial"
        r2.font.size = Pt(8)
        r2.font.bold = True
        r2.font.color.rgb = RGB_BLUE

    doc.save(output_path)
    print(f"[+] Professionally styled DOCX saved: {output_path}")


def build_pdf_report(data, output_path):
    # Sanitize all strings with html.escape to prevent unescaped HTML tags (e.g. <h1>, <img>, < 20 Words)
    # from being rendered as actual DOM elements!
    issues_html = ""
    for idx, iss in enumerate(data.get("issues", []), 1):
        title_esc = html.escape(iss.get('title', ''))
        urls_esc = html.escape(iss.get('urls', ''))
        impact_esc = html.escape(iss.get('impact', ''))
        issues_html += f"""
        <div class="callout">
            <h4>{idx}. {title_esc}</h4>
            <p><strong>Affected URLs:</strong> <code>{urls_esc}</code><br>
            <strong>Observed:</strong> {impact_esc}</p>
        </div>
        """

    scorecard_html = ""
    for r in data.get("scorecard", []):
        badge_cls = "badge-critical" if r[3] in ["Fail", "Critical"] else ("badge-warning" if r[3] == "Warning" else "badge-pass")
        scorecard_html += f"""
        <tr>
            <td><strong>{html.escape(str(r[0]))}</strong></td>
            <td>{html.escape(str(r[1]))}</td>
            <td><strong>{html.escape(str(r[2]))}</strong></td>
            <td><span class="badge {badge_cls}">{html.escape(str(r[3]))}</span></td>
            <td>{html.escape(str(r[4]))}</td>
        </tr>
        """

    scope_html = ""
    for sc in data.get("scope_of_improvement", []):
        pfx = html.escape(sc.get('bold_prefix', ''))
        txt = html.escape(sc.get('text', ''))
        scope_html += f"<p><strong>{pfx}</strong> {txt}</p>\n"

    fixes_html = ""
    for idx, fx in enumerate(data.get("immediate_fixations", []), 1):
        fx_esc = html.escape(fx)
        if ":" in fx_esc:
            prefix, rest = fx_esc.split(":", 1)
            fixes_html += f"<li><strong>{idx}. {prefix}:</strong>{rest}</li>\n"
        else:
            fixes_html += f"<li><strong>{idx}.</strong> {fx_esc}</li>\n"

    roadmap_html = ""
    for rm in data.get("roadmap", []):
        ph_esc = html.escape(rm.get('phase', '')).replace('\n', '<br>')
        obj_esc = html.escape(rm.get('objectives', ''))
        ms_esc = html.escape(rm.get('milestone', ''))
        roadmap_html += f"""
        <tr>
            <td><strong>{ph_esc}</strong></td>
            <td>{obj_esc}</td>
            <td>{ms_esc}</td>
        </tr>
        """

    exec_html = ""
    for ex in data.get("executive_summary", []):
        pfx = html.escape(ex.get('bold_prefix', ''))
        txt = html.escape(ex.get('text', ''))
        exec_html += f"<p><strong>{pfx}</strong> {txt}</p>\n"

    report_title_esc = html.escape(data.get('report_title', 'SEO & Website Audit Report'))
    client_name_esc = html.escape(data.get('client_name', 'Client'))
    badge_text_esc = html.escape(data.get('badge_text', 'B2B Technical SEO & On-Page Audit'))
    domain_esc = html.escape(data.get('domain', ''))
    primary_svc_esc = html.escape(data.get('primary_service', ''))
    core_offering_esc = html.escape(data.get('core_offering', ''))
    tech_stack_esc = html.escape(data.get('technology_stack', ''))
    audit_date_esc = html.escape(data.get('audit_date', ''))
    da_summary_esc = html.escape(data.get('domain_authority_summary', ''))
    health_score_esc = html.escape(str(data.get('health_score', 50)))

    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>{report_title_esc}</title>
    <style>
        @page {{
            size: A4;
            margin: 12mm 12mm 12mm 12mm;
            @bottom-center {{
                content: "{client_name_esc} - SEO Audit | Page " counter(page);
                font-size: 8pt;
                color: #64748b;
            }}
        }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
            color: #0f172a;
            line-height: 1.4;
            font-size: 8.5pt;
            background: #ffffff;
        }}
        .cover-page {{
            text-align: center;
            padding-top: 40px;
            padding-bottom: 45px;
            page-break-after: always;
        }}
        .cover-badge {{
            display: inline-block;
            text-transform: uppercase;
            font-size: 8.5pt;
            font-weight: 700;
            color: #2563eb;
            letter-spacing: 1px;
            margin-bottom: 10px;
        }}
        .cover-title {{
            font-size: 22pt;
            font-weight: 800;
            color: #0a192f;
            margin-bottom: 8px;
            line-height: 1.2;
        }}
        .cover-subtitle {{
            font-size: 11.5pt;
            color: #475569;
            margin-bottom: 30px;
        }}
        .score-box {{
            display: inline-block;
            border: 3px solid #dc2626;
            background: #fef2f2;
            border-radius: 12px;
            padding: 20px 36px;
            margin-bottom: 30px;
        }}
        .score-val {{
            font-size: 42pt;
            font-weight: 800;
            color: #dc2626;
            line-height: 1;
        }}
        .score-lbl {{
            font-size: 9pt;
            font-weight: 700;
            color: #991b1b;
            text-transform: uppercase;
            letter-spacing: 0.5px;
            margin-top: 5px;
        }}
        .cover-meta {{
            max-width: 520px;
            margin: 0 auto;
            text-align: left;
            background: #f8fafc;
            border: 1px solid #e2e8f0;
            border-radius: 8px;
            padding: 14px 18px;
            font-size: 8.5pt;
        }}
        .cover-meta div {{ margin-bottom: 4px; }}
        .cover-meta strong {{ color: #0a192f; }}
        .section {{ margin-bottom: 16px; }}
        .section-title {{
            font-size: 11pt;
            font-weight: 800;
            color: #0a192f;
            border-bottom: 2px solid #2563eb;
            padding-bottom: 3px;
            margin-bottom: 6px;
            margin-top: 10px;
        }}
        p {{ margin-bottom: 5px; color: #334155; text-align: justify; }}
        table {{ width: 100%; border-collapse: collapse; font-size: 7.5pt; margin: 6px 0 12px; page-break-inside: avoid; }}
        th {{
            background: #0a192f;
            color: #ffffff;
            font-weight: 700;
            text-align: left;
            padding: 5px 8px;
            text-transform: uppercase;
            font-size: 7pt;
            letter-spacing: 0.5px;
        }}
        td {{ padding: 5px 8px; border-bottom: 1px solid #e2e8f0; color: #334155; vertical-align: top; }}
        tr:nth-child(even) td {{ background: #f8fafc; }}
        .badge {{ display: inline-block; padding: 1px 4px; border-radius: 3px; font-size: 6.5pt; font-weight: 700; text-transform: uppercase; }}
        .badge-critical {{ background: #fee2e2; color: #dc2626; }}
        .badge-warning {{ background: #fef3c7; color: #d97706; }}
        .badge-pass {{ background: #d1fae5; color: #059669; }}
        .callout {{
            background: #fff5f5;
            border-left: 3px solid #dc2626;
            padding: 7px 11px;
            border-radius: 0 5px 5px 0;
            margin-bottom: 6px;
            page-break-inside: avoid;
        }}
        .callout h4 {{ font-size: 8.5pt; font-weight: 700; color: #991b1b; margin-bottom: 3px; }}
        .callout p {{ font-size: 7.5pt; color: #334155; margin-bottom: 0; line-height: 1.45; text-align: left; }}
        .callout code {{
            background: #f1f5f9;
            color: #0f172a;
            padding: 1px 4px;
            border-radius: 3px;
            font-size: 7pt;
            word-break: break-all;
        }}
        ul {{ margin-left: 18px; font-size: 8pt; color: #334155; line-height: 1.5; }}
        li {{ margin-bottom: 4px; }}
    </style>
</head>
<body>

<div class="cover-page">
    <div class="cover-badge">{badge_text_esc}</div>
    <div class="cover-title">{report_title_esc}</div>
    <div class="cover-subtitle">{client_name_esc} ({domain_esc}) — Issues &amp; Gap Analysis</div>

    <div class="score-box">
        <div class="score-val">{health_score_esc} / 100</div>
        <div class="score-lbl">Overall Website Health Score</div>
    </div>

    <div class="cover-meta">
        <div><strong>Client Domain:</strong> {domain_esc}</div>
        <div><strong>Primary Service / Focus:</strong> {primary_svc_esc}</div>
        <div><strong>Core Offering:</strong> {core_offering_esc}</div>
        <div><strong>Technology Stack:</strong> {tech_stack_esc}</div>
        <div><strong>Audit Verification:</strong> {audit_date_esc}</div>
        <div><strong>Domain Authority:</strong> {da_summary_esc}</div>
    </div>
</div>

<div class="section">
    <div class="section-title">1. Executive Summary &amp; Live-Verified Baseline</div>
    {exec_html}
</div>

<div class="section">
    <div class="section-title">2. Website Technical Health Scorecard ({health_score_esc} / 100)</div>
    <table>
        <thead>
            <tr><th>Audit Category</th><th>Weight</th><th>Score</th><th>Status</th><th>Primary Observed Defect</th></tr>
        </thead>
        <tbody>
            {scorecard_html}
        </tbody>
    </table>
</div>

<div class="section" style="page-break-before: always;">
    <div class="section-title">3. Comprehensive Issue Catalog ({len(data.get('issues', []))} Verified Issues - Issues ONLY)</div>
    {issues_html}
</div>

<div class="section">
    <div class="section-title">4. Scope of Improvement</div>
    {scope_html}
</div>

<div class="section">
    <div class="section-title">5. Immediate Fixation Needed (Emergency Priority - {len(data.get('immediate_fixations', []))} Items)</div>
    <ul>
        {fixes_html}
    </ul>
</div>

<div class="section">
    <div class="section-title">6. What Will Be the Next Step (Strategic Phased Roadmap)</div>
    <table>
        <thead>
            <tr><th>Phase</th><th>Key Objectives &amp; Focus Areas</th><th>Target Milestone</th></tr>
        </thead>
        <tbody>
            {roadmap_html}
        </tbody>
    </table>
</div>

</body>
</html>
"""

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        page.set_content(html_content)
        page.wait_for_load_state("networkidle")
        page.pdf(path=output_path, format="A4", print_background=True, margin={"top": "12mm", "bottom": "12mm", "left": "12mm", "right": "12mm"})
        browser.close()

    print(f"[+] PDF Report saved: {output_path} ({os.path.getsize(output_path):,} bytes)")


def main():
    parser = argparse.ArgumentParser(description="Universal Issue-Only SEO Audit Report Generator (DOCX & PDF)")
    parser.add_argument("--data", required=True, help="Path to input audit JSON data file")
    parser.add_argument("--out", required=True, help="Directory to save generated DOCX and PDF files")
    parser.add_argument("--prefix", default="Client_SEO_Audit_Report", help="Output filename prefix (without extension)")
    args = parser.parse_args()

    os.makedirs(args.out, exist_ok=True)
    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    docx_path = os.path.join(args.out, f"{args.prefix}.docx")
    pdf_path = os.path.join(args.out, f"{args.prefix}.pdf")

    build_docx_report(data, docx_path)
    build_pdf_report(data, pdf_path)
    print("\n[+] Both executive reports generated successfully!")


if __name__ == "__main__":
    main()
